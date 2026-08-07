"""
Generación de reportes Excel. Reglas:
- openpyxl directo, sin pandas -- no necesitamos DataFrames intermedios para
  un volcado de columnas, y evita una dependencia pesada solo para esto.
- El filtro de ownership (quién ve qué) NO se decide aquí -- se reutiliza
  ShipmentService.list_shipments(), que ya aplica la regla real. Un reporte
  nunca debe ser una puerta trasera para ver datos que la GUI no muestra.
"""
from __future__ import annotations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADERS = [
    "Canal", "OV", "Cliente", "Estatus", "Motivo",
    "Cajas", "Bolsas", "Ubicación destino", "Paquetería", "Valor (MXN)",
    "Número de guía", "Horario entrega", "Quién entrega", "Chofer recibe",
    "Hora salida", "Días de estancia", "Última actualización",
]

_HEADER_FILL = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _row_from_shipment(s) -> list:
    return [
        s.channel, s.order_number, s.customer_name, s.status, s.problem_type or "",
        s.cajas, s.bolsas, s.ubicacion_destino, s.paqueteria, s.valor_mxn,
        s.numero_guia, s.horario_entrega, s.nombre_quien_entrega, s.nombre_chofer_recibe,
        s.hora_salida, s.dias_estancia,
        s.updated_at.strftime("%Y-%m-%d %H:%M") if s.updated_at else "",
    ]


def _write_sheet(ws, shipments: list) -> None:
    ws.append(_HEADERS)
    for col_idx in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for s in shipments:
        ws.append(_row_from_shipment(s))

    # Autoajuste simple de ancho de columna -- no es exacto pixel-perfect,
    # pero evita que el reporte salga con columnas de 8 caracteres ilegibles.
    for col_idx, header in enumerate(_HEADERS, start=1):
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            value = row[0].value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)


def build_consolidated_workbook(shipments: list, output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    _write_sheet(ws, shipments)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def build_channel_workbook(shipments: list, channel_code: str, output_path: str) -> None:
    filtered = [s for s in shipments if s.channel == channel_code]
    wb = Workbook()
    ws = wb.active
    ws.title = channel_code[:31]  # límite de Excel para nombres de hoja
    _write_sheet(ws, filtered)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
